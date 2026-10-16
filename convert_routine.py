import openpyxl
import json
import re
import os

file_path = "routine.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=False)
routine_data = []

sheet_day_map = {
    'Wed - Upper A': {'id': 'wed___upper_a', 'title': 'Wed - Upper A', 'tag': 'Upper'},
    'Fri - Lower A': {'id': 'fri___lower_a', 'title': 'Fri - Lower A', 'tag': 'Lower'},
    'Sat - Upper B': {'id': 'sat___upper_b', 'title': 'Sat - Upper B', 'tag': 'Upper'},
    'Mon- Lower B': {'id': 'mon__lower_b', 'title': 'Mon - Lower B', 'tag': 'Lower'},
    'Mon - Lower B': {'id': 'mon__lower_b', 'title': 'Mon - Lower B', 'tag': 'Lower'},
    'Sun - DB Arms & Shoulders': {'id': 'sun___db_arms', 'title': 'Sun - DB Arms & Shoulders', 'tag': 'Arms / Shoulders'}
}

# Load existing routine-data.js metadata before writing
existing_routine = []
existing_metadata = {}
trailing_js = ""

if os.path.exists('routine-data.js'):
    with open('routine-data.js', 'r', encoding='utf-8') as f:
        existing_js = f.read()
    
    m = re.search(r'const DEFAULT_ROUTINE = (\[[\s\S]*?\]);\s*(?:\n([\s\S]*))?$', existing_js)
    if m:
        try:
            existing_routine = json.loads(m.group(1))
            trailing_js = m.group(2) or ""
        except Exception as e:
            print("Warning: could not parse existing DEFAULT_ROUTINE JSON:", e)
    
    for day in existing_routine:
        for ex in day.get('exercises', []):
            name_key = ex.get('name')
            if name_key:
                existing_metadata[(day.get('id'), name_key)] = ex
                if name_key not in existing_metadata:
                    existing_metadata[name_key] = ex

def clean_text(s):
    if s is None:
        return ''
    s = str(s).strip()
    # Replace the replacement char with proper punctuation
    # 4560 -> 45°-60°, 45 -> 45°, 30 -> 30°
    s = s.replace('45\ufffd60\ufffd', '45°–60°')
    s = s.replace('45\ufffd', '45°')
    s = s.replace('30\ufffd', '30°')
    s = s.replace('\ufffd', '–')
    return s.strip()

for sname in wb.sheetnames:
    ws = wb[sname]
    
    if sname in sheet_day_map:
        day_info = sheet_day_map[sname]
        day_id = day_info['id']
        day_title = day_info['title']
        day_tag = day_info['tag']
    else:
        day_id = re.sub(r'[^a-zA-Z0-9]', '_', sname.lower()).strip('_')
        day_title = sname
        day_tag = 'Upper' if 'Upper' in sname else 'Lower'
    
    day_obj = {
        'id': day_id,
        'title': day_title,
        'tag': day_tag,
        'exercises': []
    }
    
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if not name:
            continue
        sets_val = ws.cell(r, 2).value
        try:
            sets = int(float(sets_val))
        except (ValueError, TypeError):
            sets = 3
            
        reps = clean_text(ws.cell(r, 3).value)
        weight = clean_text(ws.cell(r, 4).value)
        if weight == '140.0':
            weight = '140 lbs'
            
        rest = clean_text(ws.cell(r, 5).value)
        notes = clean_text(ws.cell(r, 6).value)
        
        # Hyperlink formula in column 7
        video_cell = ws.cell(r, 7).value
        video_url = ''
        if video_cell and 'HYPERLINK' in str(video_cell):
            m = re.search(r'HYPERLINK\("([^"]+)"', str(video_cell))
            if m:
                video_url = m.group(1)
        
        # Clean up exercise name
        clean_name = clean_text(name)
        lower_name = clean_name.lower()
        
        # All exercises are core exercises
        is_optional = False
            
        # Target muscle category
        muscle = 'Full Body'
        if any(w in lower_name for w in ['lateral raise', 'shoulder press', 'overhead shoulder']):
            muscle = 'Shoulders'
        elif any(w in lower_name for w in ['chest press', 'pec deck', 'incline press']):
            muscle = 'Chest'
        elif any(w in lower_name for w in ['pulldown', 'row', 'lat pulldown']):
            muscle = 'Back / Lats'
        elif any(w in lower_name for w in ['calf']):
            muscle = 'Calves'
        elif any(w in lower_name for w in ['hack squat', 'leg extension', 'leg press']):
            muscle = 'Quads'
        elif any(w in lower_name for w in ['leg curl', 'rdl', 'deadlift', 'back extension']):
            muscle = 'Hamstrings / Glutes'
        elif any(w in lower_name for w in ['hip thrust', 'glute drive']):
            muscle = 'Glutes'
        elif any(w in lower_name for w in ['adduction', 'abduction']):
            muscle = 'Adductors / Hips'
        elif any(w in lower_name for w in ['crunch', 'knee raise', 'core']):
            muscle = 'Abs / Core'
        elif any(w in lower_name for w in ['tricep', 'pushdown', 'overhead cable']):
            muscle = 'Triceps'
        elif any(w in lower_name for w in ['bicep', 'bayonet']):
            muscle = 'Biceps'
        elif 'curl' in lower_name and 'leg' not in lower_name:
            muscle = 'Biceps'
            
        # Superset pairings logic
        superset = None
        if 'wed' in day_id:
            if 'triceps rope' in lower_name:
                superset = {
                    'id': 'upper_a_arms',
                    'tag': '5A',
                    'name': 'Arm Finisher Superset (Triceps + Biceps)',
                    'tip': '⚡ Superset: Perform 1 set of Triceps, rest 45–60s, then perform 1 set of Biceps, rest 60s. Saves ~6 mins!',
                    'rest': '45–60 sec'
                }
            elif 'cable biceps' in lower_name:
                superset = {
                    'id': 'upper_a_arms',
                    'tag': '5B',
                    'name': 'Arm Finisher Superset (Triceps + Biceps)',
                    'tip': '⚡ Superset: Alternate with Triceps Rope Pushdown. Saves ~6 mins!',
                    'rest': '60 sec'
                }
        elif 'fri' in day_id:
            if 'leg extension' in lower_name:
                superset = {
                    'id': 'lower_a_quad_ham',
                    'tag': '2A',
                    'name': 'Antagonist Superset (Quads + Hamstrings)',
                    'tip': '⚡ Superset: Perform 1 set of Leg Extensions, rest 60s, then perform 1 set of Seated Leg Curls, rest 60s. Saves ~8 mins!',
                    'rest': '60 sec'
                }
            elif 'seated leg curl' in lower_name:
                superset = {
                    'id': 'lower_a_quad_ham',
                    'tag': '2B',
                    'name': 'Antagonist Superset (Quads + Hamstrings)',
                    'tip': '⚡ Superset: Alternate with Leg Extensions. Quads rest while hamstrings work!',
                    'rest': '60–75 sec'
                }
        elif 'sat' in day_id:
            if 'overhead cable triceps' in lower_name:
                superset = {
                    'id': 'upper_b_arms',
                    'tag': '5A',
                    'name': 'Cable Arm Superset (Long-head Triceps + Lengthened Biceps)',
                    'tip': '⚡ Superset: Same cable station! Do Overhead Extension, rest 45s, do Bayonet Curl, rest 60s. Saves ~6 mins!',
                    'rest': '45–60 sec'
                }
            elif 'bayonet cable curl' in lower_name:
                superset = {
                    'id': 'upper_b_arms',
                    'tag': '5B',
                    'name': 'Cable Arm Superset (Long-head Triceps + Lengthened Biceps)',
                    'tip': '⚡ Superset: Alternate with Overhead Triceps Extensions. Saves ~6 mins!',
                    'rest': '60 sec'
                }
        elif 'mon' in day_id:
            if 'machine hip adduction' in lower_name:
                superset = {
                    'id': 'lower_b_acc',
                    'tag': '4A',
                    'name': 'Accessory Superset (Adductors + Calves)',
                    'tip': '⚡ Superset: Perform 1 set of Adduction, rest 45s, perform Calf Press, rest 60s. Saves ~6 mins!',
                    'rest': '45–60 sec'
                }
            elif 'leg press calf press' in lower_name:
                superset = {
                    'id': 'lower_b_acc',
                    'tag': '4B',
                    'name': 'Accessory Superset (Adductors + Calves)',
                    'tip': '⚡ Superset: Alternate with Hip Adduction. Saves ~6 mins!',
                    'rest': '60 sec'
                }

        # Check existing metadata preservation
        meta = existing_metadata.get((day_id, clean_name)) or existing_metadata.get(clean_name)
        ex_id = meta.get('id') if meta and meta.get('id') else f'{day_id}_ex{r-1}'
        if meta and meta.get('muscle'):
            muscle = meta.get('muscle')
        if meta and meta.get('superset') is not None:
            superset = meta.get('superset')
        if meta and meta.get('sets') is not None:
            sets = meta.get('sets')
        if meta and meta.get('targetReps'):
            reps = meta.get('targetReps')
        if meta and meta.get('notes'):
            notes = meta.get('notes')
        if meta and meta.get('startingWeight'):
            weight = meta.get('startingWeight')
        if meta and meta.get('rest'):
            rest = meta.get('rest')
        if meta and meta.get('videoUrl'):
            video_url = meta.get('videoUrl')
        
        # Calculate or preserve increment (+10 for compounds, +5 for isolation)
        isolation_keywords = [
            'dumbbell', 'db ', 'single-arm', 'single-leg', 'one-arm', 'lateral raise',
            'rear delt', 'face pull', 'pec deck', 'fly', 'crossover', 'leg extension',
            'leg curl', 'sissy squat', 'goblet squat', 'split squat', 'curl', 'tricep',
            'pushdown', 'skull crusher', 'extension', 'crunch', 'knee raise', 'sit-up',
            'pull-through', 'pullover', 'adduction', 'abduction'
        ]
        compound_keywords = [
            'hack squat', 'leg press', 'smith machine squat', 'hip thrust', 'glute drive',
            'barbell romanian', 'smith machine romanian', 'calf raise', 'calf press',
            'machine chest press', 'machine incline press', 'smith machine incline press',
            'smith machine chest press', 'machine overhead', 'machine shoulder press',
            'smith machine overhead', 'smith machine shoulder', 'lat pulldown',
            'machine row', 'seated cable row', 't-bar row'
        ]
        if meta and meta.get('increment') is not None:
            increment = meta.get('increment')
        elif any(kw in lower_name for kw in isolation_keywords):
            increment = 5
        elif any(kw in lower_name for kw in compound_keywords):
            increment = 10
        else:
            increment = 5

        ex_obj = {
            'id': ex_id,
            'name': clean_name,
            'muscle': muscle,
            'sets': sets,
            'targetReps': reps,
            'startingWeight': weight,
            'rest': rest,
            'notes': notes,
            'videoUrl': video_url,
            'isOptional': False,
            'superset': superset,
            'increment': increment
        }
        if meta and 'alternatives' in meta:
            ex_obj['alternatives'] = meta['alternatives']

        day_obj['exercises'].append(ex_obj)

    # Preserve customized day exercise list from existing routine-data.js if present
    existing_day = next((d for d in existing_routine if d.get('id') == day_id), None)
    if existing_day and existing_day.get('exercises'):
        preserved_exercises = []
        for ex in existing_day['exercises']:
            ex_copy = dict(ex)
            ex_copy['isOptional'] = False
            if 'increment' not in ex_copy:
                lname = ex_copy.get('name', '').lower()
                if any(kw in lname for kw in isolation_keywords):
                    ex_copy['increment'] = 5
                elif any(kw in lname for kw in compound_keywords):
                    ex_copy['increment'] = 10
                else:
                    ex_copy['increment'] = 5
            preserved_exercises.append(ex_copy)
        day_obj['exercises'] = preserved_exercises
    elif 'wed' in day_id and not any('opt_facepull' in e.get('id', '') or 'Face Pull' in e.get('name', '') for e in day_obj['exercises']):
        meta = existing_metadata.get((day_id, 'Cable Face Pull (Rear Delts & Posture)'))
        day_obj['exercises'].append({
            'id': meta.get('id') if meta else f'{day_id}_opt_facepull',
            'name': 'Cable Face Pull (Rear Delts & Posture)',
            'muscle': 'Shoulders',
            'sets': 2,
            'targetReps': '10–12',
            'startingWeight': '30–40 lbs',
            'rest': '90 sec',
            'notes': 'Set pulley to eye level; thumbs back; pull rope toward forehead while pulling elbows back to build healthy shoulders.',
            'videoUrl': 'https://www.youtube.com/results?search_query=Jeff+Nippard+face+pull+technique',
            'isOptional': False,
            'superset': None,
            'increment': 5
        })

    routine_data.append(day_obj)

# Preserve Sunday DB Arms & Shoulders from existing routine if not present in Excel workbook
existing_sunday = next((d for d in existing_routine if d.get('id') == 'sun___db_arms'), None)
if existing_sunday and not any(d.get('id') == 'sun___db_arms' for d in routine_data):
    for ex in existing_sunday.get('exercises', []):
        ex['isOptional'] = False
    mon_idx = next((i for i, d in enumerate(routine_data) if 'mon' in d.get('id', '')), -1)
    if mon_idx != -1:
        routine_data.insert(mon_idx, existing_sunday)
    else:
        routine_data.append(existing_sunday)

with open('routine-data.js', 'w', encoding='utf-8') as f:
    f.write('// Pre-loaded routine parsed from routine.xlsx + Sunday DB Arms & Shoulders\n')
    f.write('const DEFAULT_ROUTINE = ')
    f.write(json.dumps(routine_data, indent=2, ensure_ascii=False))
    f.write(';\n\n')
    if trailing_js.strip():
        f.write(trailing_js.strip() + '\n')
    else:
        f.write('if (typeof module !== "undefined") { module.exports = { DEFAULT_ROUTINE }; }\n')

print(f"Generated routine-data.js with {len(routine_data)} workout days.")
for d in routine_data:
    print(f" - {d['title']}: {len(d['exercises'])} exercises (Optional: {[e['name'] for e in d['exercises'] if e.get('isOptional')]})")
